from __future__ import annotations

import hashlib
import math
import re
from pathlib import Path
from typing import Any

from sqlalchemy import func, select, update
from sqlalchemy.orm import Session

from ...config import UPLOAD_DIR
from ...models import ExtractedItem, IngestionJob, SourceDocument
from ..math_markup import split_intact


FORMULA_HINT = re.compile(r"(?:[A-Za-z][A-Za-z0-9_{}]*\s*[=<>≥≤]\s*[^。！？\n]{1,100})")
HEADING_HINT = re.compile(r"^(?:第\s*[一二三四五六七八九十百0-9]+\s*[章节]|[0-9]+(?:\.[0-9]+)*\s+).{2,80}$")
_SAFE_NAME = re.compile(r"[^A-Za-z0-9_.\-\u3400-\u9fff]+")
_OCR_ENGINE: Any | None = None


class OCRUnavailable(RuntimeError):
    """OCR cannot run because its optional CPU runtime is unavailable."""


def sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _classify(text: str) -> str:
    if FORMULA_HINT.search(text):
        return "formula"
    if "例" in text[:40] or "解：" in text or "解答" in text[:80]:
        return "example"
    if "图" in text[:40] or "波形" in text or "频谱" in text:
        return "waveform"
    return "text"


def _safe_json(value: Any, *, depth: int = 0) -> Any:
    """Keep model metadata JSON-safe; never evaluate formulas or numeric strings."""
    if depth > 5:
        return None
    if value is None or isinstance(value, (str, bool)):
        return value[:4000] if isinstance(value, str) else value
    if isinstance(value, int):
        return value if abs(value) <= 10**12 else None
    if isinstance(value, float):
        return value if math.isfinite(value) and abs(value) <= 10**12 else None
    if isinstance(value, list):
        return [_safe_json(item, depth=depth + 1) for item in value[:200]]
    if isinstance(value, dict):
        return {str(key)[:100]: _safe_json(item, depth=depth + 1) for key, item in list(value.items())[:200]}
    return str(value)[:4000]


def _page_bounds(page: Any) -> list[float]:
    rect = page.rect
    return [float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)]


def _text_blocks(page: Any) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    try:
        data = page.get_text("dict")
        for block in data.get("blocks", []):
            lines: list[str] = []
            for line in block.get("lines", []):
                lines.append("".join(str(span.get("text", "")) for span in line.get("spans", [])))
            text = "\n".join(lines).strip()
            if text:
                blocks.append({"bbox": [float(v) for v in block.get("bbox", [])], "text": text})
    except (AttributeError, TypeError, ValueError):
        return []
    return blocks

def _page_image_coverage(page: Any) -> float:
    """Return the largest image coverage ratio, without trusting image text."""
    try:
        rect = page.rect
        page_area = float(rect.width) * float(rect.height)
        if page_area <= 0:
            return 0.0
        largest = 0.0
        for image in page.get_image_info(xrefs=True):
            bbox = image.get("bbox") if isinstance(image, dict) else None
            if not isinstance(bbox, (list, tuple)) or len(bbox) != 4:
                continue
            x0, y0, x1, y1 = (float(value) for value in bbox)
            largest = max(largest, max(0.0, x1 - x0) * max(0.0, y1 - y0))
        return min(1.0, largest / page_area)
    except (AttributeError, TypeError, ValueError, ZeroDivisionError):
        return 0.0


def _pdf_quality(page: Any, raw_text: str) -> dict[str, Any]:
    image_count = len(page.get_images(full=True))
    image_coverage = _page_image_coverage(page)
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    short_line_ratio = (
        sum(len(line) <= 3 for line in lines) / len(lines)
        if lines else 0.0
    )
    quality_flags: list[str] = []
    if len(raw_text) < 40:
        quality_flags.append("ocr_required")
    if raw_text and image_coverage >= 0.8:
        quality_flags.append("full_page_image_native_text")
    if raw_text and image_count and len(lines) >= 20 and short_line_ratio >= 0.3:
        quality_flags.append("fragmented_native_text")
    return {
        "images": image_count,
        "image_coverage": round(image_coverage, 4),
        "short_line_ratio": round(short_line_ratio, 4),
        "quality_flags": quality_flags,
        "text_quality": "needs_review" if quality_flags else "native",
    }


def _quality_flags(payload: dict[str, Any]) -> list[str]:
    value = payload.get("quality_flags", [])
    flags = [flag for flag in value if isinstance(flag, str) and flag] if isinstance(value, list) else []
    if not flags and payload.get("text_quality") == "needs_review":
        flags.append("low_quality_text")
    return flags


def _quality_review_cleared(payload: dict[str, Any]) -> bool:
    return payload.get("quality_review_cleared") is True


def _add_quality_flag(payload: dict[str, Any], flag: str) -> list[str]:
    flags = _quality_flags(payload)
    if flag not in flags:
        flags.append(flag)
    payload["quality_flags"] = flags
    payload["text_quality"] = "needs_review"
    return flags


def _ensure_legacy_quality_flags(item: ExtractedItem, document: SourceDocument) -> dict[str, Any]:
    """Annotate old native-PDF rows that predate the quality gate."""
    payload = dict(item.payload or {})
    if document.parser == "pymupdf" and payload.get("images") and "image_coverage" not in payload and not _quality_flags(payload):
        _add_quality_flag(payload, "legacy_image_backed_native_text")
    return payload


def _revoke_material(material: Any, reason: str) -> None:
    metadata = dict(material.metadata_json or {})
    metadata["publication_status"] = "revoked"
    metadata["revoked_reason"] = reason
    material.metadata_json = metadata
    material.verified = False


def _material_publication_status(material: Any) -> str:
    return str((material.metadata_json or {}).get("publication_status", "published"))


def _invalidate_item_publication(db: Session, item: ExtractedItem, reason: str) -> None:
    """Keep historical IDs, but exclude a source awaiting re-review from retrieval."""
    from ...models import Material

    job = db.get(IngestionJob, item.ingestion_job_id)
    document_id = job.source_document_id if job else None
    if document_id is None:
        return
    material = db.scalar(
        select(Material).where(
            Material.source_document_id == document_id,
            Material.page_number == item.page_number,
            Material.source_anchor == f"doc:{document_id}:p{item.page_number}",
        )
    )
    if material:
        _revoke_material(material, reason)


def render_page_bytes(source: SourceDocument | Path | str, page_number: int, *, dpi: int = 144) -> bytes:
    """Render one original PDF page to PNG bytes without modifying the source file."""
    import fitz

    path = Path(source.path if isinstance(source, SourceDocument) else source)
    if path.suffix.lower() != ".pdf":
        raise ValueError("只有 PDF 支持页面渲染")
    if page_number < 1:
        raise ValueError("页码必须从 1 开始")
    with fitz.open(path) as document:
        if page_number > len(document):
            raise ValueError("页码超出原始 PDF 范围")
        scale = max(0.5, min(float(dpi) / 72.0, 8.0))
        pixmap = document[page_number - 1].get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        return pixmap.tobytes("png")




def _extract_pdf(path: Path) -> tuple[int, list[dict[str, Any]], list[dict[str, Any]], str]:
    import fitz

    items: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    current_section = ""
    with fitz.open(path) as doc:
        for index, page in enumerate(doc, start=1):
            raw_text = page.get_text("text").strip()
            blocks = _text_blocks(page)
            quality = _pdf_quality(page, raw_text)
            image_count = quality["images"]
            image_coverage = quality["image_coverage"]
            lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
            short_line_ratio = quality["short_line_ratio"]
            payload = {
                "section": current_section,
                "formulas": FORMULA_HINT.findall(raw_text),
                "images": image_count,
                "image_coverage": round(image_coverage, 4),
                "short_line_ratio": round(short_line_ratio, 4),
                "source_bounds": _page_bounds(page),
                "page_bounds": _page_bounds(page),
                "text_blocks": blocks,
            }
            quality_flags: list[str] = quality["quality_flags"]
            if quality_flags:
                payload["quality_flags"] = quality_flags
                payload["text_quality"] = quality["text_quality"]
                warnings.append({
                    "page": index,
                    "code": "low_quality_text",
                    "flags": quality_flags,
                    "message": "页面原生文本不可直接信任，需要显式 OCR 或人工核对原页",
                })
            if "ocr_required" in quality_flags:
                payload["needs_ocr"] = True
                items.append({
                    "page_number": index, "item_type": "page", "raw_text": raw_text,
                    "normalized_text": raw_text, "payload": payload, "confidence": 0.0,
                    "needs_review": True, "review_status": "pending", "source_anchor": "",
                })
                continue
            lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
            headings = [line for line in lines if HEADING_HINT.match(line)]
            if headings:
                current_section = headings[0]
                payload["section"] = current_section
            formulas = payload["formulas"]
            item_type = _classify(raw_text)
            base_confidence = 0.72 if len(raw_text) > 200 else 0.55
            confidence = min(base_confidence, 0.35) if quality_flags else base_confidence
            items.append({
                "page_number": index, "item_type": item_type, "raw_text": raw_text,
                "normalized_text": raw_text, "payload": payload, "confidence": confidence,
                "needs_review": bool(quality_flags) or item_type in {"formula", "waveform", "spectrum"} or bool(formulas),
                "review_status": "pending", "source_anchor": "",
            })
        return len(doc), items, warnings, "pymupdf"

def _extract_docx(path: Path) -> tuple[int, list[dict[str, Any]], list[dict[str, Any]], str]:
    from docx import Document

    document = Document(path)
    items: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        item_type = "text"
        if paragraph.style and "Heading" in paragraph.style.name:
            item_type = "heading"
        elif FORMULA_HINT.search(text):
            item_type = "formula"
        items.append({
            "page_number": index, "item_type": item_type, "raw_text": text,
            "normalized_text": text, "payload": {"style": paragraph.style.name if paragraph.style else "", "formulas": FORMULA_HINT.findall(text)},
            "confidence": 0.9, "needs_review": item_type == "formula", "review_status": "pending", "source_anchor": "",
        })
    return len(items), items, [], "python-docx"


def create_ingestion_job(db: Session, source_path: Path, course_id: int | None = None) -> IngestionJob:
    source_path = source_path.resolve()
    if not source_path.exists():
        raise FileNotFoundError(source_path)
    digest = sha256_path(source_path)
    existing = db.scalar(select(SourceDocument).where(SourceDocument.file_hash == digest))
    if existing:
        existing_job = db.scalar(select(IngestionJob).where(IngestionJob.source_document_id == existing.id).order_by(IngestionJob.id.desc()))
        if existing_job:
            return existing_job
    extension = source_path.suffix.lower()
    if extension == ".pdf":
        page_count, items, warnings, parser = _extract_pdf(source_path)
    elif extension == ".docx":
        page_count, items, warnings, parser = _extract_docx(source_path)
    elif extension in {".md", ".markdown", ".txt"}:
        content = source_path.read_text(encoding="utf-8", errors="replace")
        page_count, items, warnings, parser = 1, [{
            "page_number": 1, "item_type": "text", "raw_text": content, "normalized_text": content,
            "payload": {"source_bounds": [0.0, 0.0, 1.0, 1.0]}, "confidence": 0.95,
            "needs_review": True, "review_status": "pending", "source_anchor": "",
        }], [], "text"
    else:
        raise ValueError("仅支持 PDF、DOCX、Markdown 和 TXT")

    document = SourceDocument(course_id=course_id, filename=source_path.name, file_hash=digest, path=str(source_path), page_count=page_count, parser=parser, status="review", warnings=warnings)
    db.add(document)
    db.flush()
    job = IngestionJob(source_document_id=document.id, status="review", processed_pages=page_count, total_pages=page_count,
                       summary={"warnings": len(warnings), "items": len(items), "ocr_pages": sum(1 for item in items if item["payload"].get("needs_ocr"))})
    db.add(job)
    db.flush()
    for item in items:
        # The database document id, not a model-generated string, is authoritative.
        item["source_anchor"] = f"doc:{document.id}:p{item['page_number']}"
        db.add(ExtractedItem(ingestion_job_id=job.id, **item))
    db.commit()
    db.refresh(job)
    return job


def _safe_basename(filename: str) -> str:
    name = Path(filename or "upload.bin").name
    name = _SAFE_NAME.sub("_", name).strip(" .") or "upload.bin"
    return name[:180]


def save_upload(filename: str, content: bytes) -> Path:
    """Store immutable uploads under content hash plus a sanitized basename."""
    digest = hashlib.sha256(content).hexdigest()
    path = UPLOAD_DIR / f"{digest}-{_safe_basename(filename)}"
    if not path.exists():
        path.write_bytes(content)
    return path


def list_job_items(db: Session, job_id: int) -> list[ExtractedItem]:
    return list(db.scalars(select(ExtractedItem).where(ExtractedItem.ingestion_job_id == job_id).order_by(ExtractedItem.page_number, ExtractedItem.id)))


def audit_job_quality(db: Session, job_id: int) -> dict[str, int | str]:
    """Re-read native PDF pages and quarantine legacy items with low-quality text."""
    import fitz

    from ...models import Material

    job = db.get(IngestionJob, job_id)
    if not job:
        raise LookupError("导入任务不存在")
    document = db.get(SourceDocument, job.source_document_id)
    if not document:
        raise LookupError("导入任务缺少原始教材")
    if document.parser != "pymupdf" or Path(document.path).suffix.lower() != ".pdf":
        return {"audited_pages": 0, "flagged": 0, "revoked": 0, "status": job.status}
    path = Path(document.path)
    if not path.exists():
        raise FileNotFoundError(path)

    audited_pages = 0
    flagged = 0
    revoked = 0
    items = list_job_items(db, job_id)
    with fitz.open(path) as pdf:
        for item in items:
            if item.page_number < 1 or item.page_number > len(pdf):
                continue
            audited_pages += 1
            page = pdf[item.page_number - 1]
            raw_text = page.get_text("text").strip()
            quality = _pdf_quality(page, raw_text)
            actual_flags = quality["quality_flags"]
            if not actual_flags:
                continue
            flagged += 1
            payload = dict(item.payload or {})
            merged_flags = list(dict.fromkeys(_quality_flags(payload) + actual_flags))
            payload.update({
                "images": quality["images"],
                "image_coverage": quality["image_coverage"],
                "short_line_ratio": quality["short_line_ratio"],
                "quality_flags": merged_flags,
                "text_quality": "needs_review",
            })
            if payload != (item.payload or {}):
                item.payload = payload
            if _quality_review_cleared(payload):
                continue
            item.needs_review = True
            if item.review_status == "approved":
                item.review_status = "pending"
            item.confidence = min(item.confidence, 0.35)
            material = db.scalar(select(Material).where(
                Material.source_document_id == document.id,
                Material.page_number == item.page_number,
                Material.source_anchor == f"doc:{document.id}:p{item.page_number}",
            ))
            if material and _material_publication_status(material) != "revoked":
                _revoke_material(material, "legacy_quality_audit_required")
                revoked += 1

    if flagged:
        job.status = "review"
        document.status = "review"
    job.summary = {
        **(job.summary or {}),
        "quality_audit_pages": audited_pages,
        "quality_audit_flagged": flagged,
        "quality_audit_revoked": revoked,
    }
    db.commit()
    return {"audited_pages": audited_pages, "flagged": flagged, "revoked": revoked, "status": job.status}


def _rag_material_metadata(document: SourceDocument, item: ExtractedItem, payload: dict[str, Any], *, verified: bool) -> dict[str, Any]:
    metadata = dict(payload)
    metadata.update({
        "rag_indexed": True,
        "rag_index_version": 2,
        "source_document_id": document.id,
        "source_page": item.page_number,
        "publication_status": "published" if verified else "indexed_unverified",
        "index_quality": "verified" if verified else "candidate",
    })
    return metadata


def _formula_sources(payload: dict[str, Any], *, verified: bool) -> list[str]:
    reviewed = payload.get("reviewed_formulas", [])
    sources = [
        row["source"].strip() for row in reviewed
        if isinstance(row, dict) and isinstance(row.get("source"), str) and row["source"].strip()
    ] if verified and isinstance(reviewed, list) else []
    if not sources:
        legacy = payload.get("formulas", [])
        sources = [value.strip() for value in legacy if isinstance(value, str) and value.strip()] if isinstance(legacy, list) else []
    return list(dict.fromkeys(sources))


def build_rag_index(db: Session, job_id: int, course_id: int | None = None) -> dict[str, int | float | str]:
    """Index every non-empty page as a traceable candidate without bypassing review."""
    from ...models import EvidenceLink, Formula, Material, MaterialChunk, Topic

    job = db.get(IngestionJob, job_id)
    if not job:
        raise LookupError("导入任务不存在")
    document = db.get(SourceDocument, job.source_document_id)
    if not document:
        raise LookupError("导入任务缺少原始教材")
    target_course_id = course_id if course_id is not None else document.course_id
    if target_course_id is None:
        raise ValueError("教材没有关联课程，无法建立 RAG 索引")
    topic_cache = {topic.name: topic for topic in db.scalars(select(Topic).where(Topic.course_id == target_course_id))}

    indexed = 0
    created = 0
    updated = 0
    unchanged = 0
    skipped_empty = 0
    skipped_rejected = 0
    verified = 0
    candidates = 0
    reviewed_empty = 0
    items = list_job_items(db, job_id)
    for item in items:
        if item.review_status == "rejected":
            skipped_rejected += 1
            continue
        payload = dict(item.payload or {})
        quality_flags = _quality_flags(payload)
        item_verified = item.review_status == "approved" and not item.needs_review and (
            not quality_flags or _quality_review_cleared(payload)
        )
        text = (item.normalized_text if item.item_type in {"blank", "image_only"}
                else item.normalized_text or item.raw_text or "").strip()
        if not text:
            skipped_empty += 1
            if item.item_type in {"blank", "image_only"}:
                reviewed_empty += int(item_verified)
                existing = db.scalar(select(Material).where(
                    Material.source_document_id == document.id, Material.page_number == item.page_number,
                    Material.source_anchor == f"doc:{document.id}:p{item.page_number}",
                ))
                if existing:
                    existing.content = ""
                    existing.verified = item_verified
                    existing.metadata_json = {
                        **_rag_material_metadata(document, item, payload, verified=item_verified),
                        "rag_indexed": False, "index_quality": item.item_type,
                    }
                    for chunk in list(existing.chunks):
                        db.execute(update(EvidenceLink).where(EvidenceLink.chunk_id == chunk.id).values(chunk_id=None))
                        db.delete(chunk)
                    for formula in list(existing.formulas):
                        db.delete(formula)
            continue
        material_type = item.item_type if item.item_type in {"text", "formula", "example", "waveform", "spectrum"} else "text"
        section = str(payload.get("topic") or payload.get("section") or "").strip()[:200]
        topic = topic_cache.get(section) if item_verified and section else None
        if item_verified and section and topic is None:
            topic = Topic(course_id=target_course_id, name=section, description="教材审核后导入")
            db.add(topic)
            db.flush()
            topic_cache[section] = topic
        source_anchor = f"doc:{document.id}:p{item.page_number}"
        metadata = _rag_material_metadata(document, item, payload, verified=item_verified)
        tags = [tag for tag in (document.filename, section) if tag]
        duplicate = db.scalar(select(Material).where(
            Material.source_document_id == document.id,
            Material.page_number == item.page_number,
            Material.source_anchor == source_anchor,
        ))
        if duplicate and (
            duplicate.content == text
            and duplicate.course_id == target_course_id
            and duplicate.material_type == material_type
            and duplicate.section == section
            and duplicate.topic_id == (topic.id if topic else None)
            and duplicate.metadata_json == metadata
            and duplicate.verified == item_verified
            and _material_publication_status(duplicate) == metadata["publication_status"]
        ):
            unchanged += 1
            indexed += 1
            verified += int(item_verified)
            candidates += int(not item_verified)
            continue
        if duplicate:
            material = duplicate
            updated += 1
            for formula in list(material.formulas):
                db.delete(formula)
        else:
            material = Material(source_document_id=document.id, source_anchor=source_anchor)
            db.add(material)
            created += 1
        material.course_id = target_course_id
        material.topic_id = topic.id if topic else None
        material.material_type = material_type
        material.title = f"第 {item.page_number} 页" + (f" · {section}" if section else "")
        material.content = text
        material.page_number = item.page_number
        material.section = section
        material.tags = tags
        material.metadata_json = metadata
        material.verified = item_verified
        db.flush()

        old_chunks = {chunk.chunk_index: chunk for chunk in material.chunks}
        for index, text_chunk in enumerate(_chunk_text(text)):
            chunk = old_chunks.pop(index, None)
            if chunk:
                chunk.text = text_chunk
                chunk.keywords = tags
            else:
                db.add(MaterialChunk(material_id=material.id, chunk_index=index, text=text_chunk, keywords=tags))
        for stale_chunk in old_chunks.values():
            db.execute(update(EvidenceLink).where(EvidenceLink.chunk_id == stale_chunk.id).values(chunk_id=None))
            db.delete(stale_chunk)
        for raw_formula in _formula_sources(payload, verified=item_verified):
            db.add(Formula(material_id=material.id, raw_formula=raw_formula, canonical_expression=raw_formula, variables_json={}))
        indexed += 1
        verified += int(item_verified)
        candidates += int(not item_verified)

    indexed_pages = {material.page_number for material in db.scalars(select(Material).where(Material.source_document_id == document.id)).all() if material.metadata_json.get("rag_indexed") and material.page_number}
    job.status = "published" if len(indexed_pages) + reviewed_empty == document.page_count and candidates == 0 else "review"
    document.status = job.status
    job.summary = {
        **(job.summary or {}),
        "rag_indexed_pages": len(indexed_pages),
        "rag_indexed_materials": indexed,
        "rag_verified_pages": verified,
        "rag_candidate_pages": candidates,
        "rag_skipped_empty": skipped_empty,
        "reviewed_empty_pages": reviewed_empty,
        "rag_skipped_rejected": skipped_rejected,
    }
    db.commit()
    coverage = round(len(indexed_pages) / document.page_count, 4) if document.page_count else 0.0
    return {
        "job_id": job.id,
        "source_document_id": document.id,
        "page_count": document.page_count,
        "indexed_pages": len(indexed_pages),
        "coverage": coverage,
        "created": created,
        "updated": updated,
        "unchanged": unchanged,
        "skipped_empty": skipped_empty,
        "skipped_rejected": skipped_rejected,
        "verified_pages": verified,
        "reviewed_empty_pages": reviewed_empty,
        "candidate_pages": candidates,
        "status": job.status,
    }


def get_rag_index_status(db: Session, job_id: int) -> dict[str, int | float | str]:
    from ...models import Material, MaterialChunk

    job = db.get(IngestionJob, job_id)
    if not job:
        raise LookupError("导入任务不存在")
    document = db.get(SourceDocument, job.source_document_id)
    if not document:
        raise LookupError("导入任务缺少原始教材")
    materials = list(db.scalars(select(Material).where(Material.source_document_id == document.id)))
    indexed_materials = [material for material in materials if (material.metadata_json or {}).get("rag_indexed") is True]
    indexed_pages = {material.page_number for material in indexed_materials if material.page_number is not None}
    verified_pages = {material.page_number for material in indexed_materials if material.page_number is not None and material.verified}
    candidate_pages = indexed_pages - verified_pages
    chunk_count = db.scalar(select(func.count(MaterialChunk.id)).where(MaterialChunk.material_id.in_([material.id for material in indexed_materials]))) if indexed_materials else 0
    reviewed_items = [item for item in list_job_items(db, job_id)
                      if item.review_status == "approved" and not item.needs_review
                      and (not _quality_flags(item.payload or {}) or _quality_review_cleared(item.payload or {}))]
    blank_pages = sum(item.item_type == "blank" for item in reviewed_items)
    image_only_pages = sum(item.item_type == "image_only" for item in reviewed_items)
    return {
        "job_id": job.id,
        "source_document_id": document.id,
        "filename": document.filename,
        "status": job.status,
        "page_count": document.page_count,
        "indexed_pages": len(indexed_pages),
        "verified_pages": len(verified_pages),
        "candidate_pages": len(candidate_pages),
        "reviewed_pages": len(reviewed_items),
        "reviewed_blank_pages": blank_pages,
        "reviewed_image_only_pages": image_only_pages,
        "review_coverage": round(len(reviewed_items) / document.page_count, 4) if document.page_count else 0.0,
        "missing_pages": max(0, document.page_count - len(indexed_pages)),
        "material_count": len(indexed_materials),
        "chunk_count": int(chunk_count or 0),
        "coverage": round(len(indexed_pages) / document.page_count, 4) if document.page_count else 0.0,
    }


def _chunk_text(text: str, max_chars: int = 1000, overlap_paragraphs: int = 1) -> list[str]:
    pieces = split_intact(text, max_chars)
    chunks: list[str] = []
    current: list[str] = []
    for piece in pieces:
        if current and len("\n\n".join(current + [piece])) > max_chars:
            chunks.append("\n\n".join(current))
            current = current[-overlap_paragraphs:] if overlap_paragraphs > 0 else []
            while current and len("\n\n".join(current + [piece])) > max_chars:
                current = current[1:]
        current.append(piece)
    if current:
        chunks.append("\n\n".join(current))
    return chunks


def _valid_bounds(payload: dict[str, Any]) -> bool:
    bounds = payload.get("source_bounds") or payload.get("page_bounds")
    return isinstance(bounds, list) and len(bounds) == 4 and all(isinstance(value, (int, float)) and math.isfinite(float(value)) for value in bounds)


def publish_job(db: Session, job_id: int, course_id: int) -> dict[str, int]:
    from ...models import EvidenceLink, Formula, Material, MaterialChunk, Topic

    job = db.get(IngestionJob, job_id)
    if not job:
        raise LookupError("导入任务不存在")
    document = db.get(SourceDocument, job.source_document_id)
    if not document:
        raise LookupError("导入任务缺少原始教材")
    items = list_job_items(db, job_id)
    topic_cache = {topic.name: topic for topic in db.scalars(select(Topic).where(Topic.course_id == course_id))}
    published = 0
    updated = 0
    skipped = 0
    revoked = 0
    for item in items:
        payload = _ensure_legacy_quality_flags(item, document)
        if payload != (item.payload or {}):
            item.payload = payload
        text = (item.normalized_text or "").strip()
        source_anchor = f"doc:{document.id}:p{item.page_number}"
        duplicate = db.scalar(select(Material).where(Material.source_document_id == document.id, Material.page_number == item.page_number, Material.source_anchor == source_anchor))
        quality_flags = _quality_flags(payload)
        if quality_flags and not _quality_review_cleared(payload):
            item.needs_review = True
            item.confidence = min(item.confidence, 0.35)
            if item.review_status == "approved":
                item.review_status = "pending"
            if duplicate and _material_publication_status(duplicate) != "revoked":
                _revoke_material(duplicate, "source_item_quality_review_required")
                revoked += 1
            skipped += 1
            continue
        risky = item.item_type in {"formula", "waveform", "spectrum"} or bool(payload.get("needs_ocr") or payload.get("ocr_engine"))
        if item.review_status != "approved" or not text or (risky and not _valid_bounds(payload)):
            if duplicate and item.review_status != "approved" and _material_publication_status(duplicate) != "revoked":
                _revoke_material(duplicate, "source_item_not_approved")
                revoked += 1
            skipped += 1
            continue
        section = str(payload.get("topic") or payload.get("section") or "").strip()[:200]
        topic = topic_cache.get(section)
        if section and topic is None:
            topic = Topic(course_id=course_id, name=section, description="教材审核后导入")
            db.add(topic)
            db.flush()
            topic_cache[section] = topic
        material_type = item.item_type if item.item_type in {"text", "formula", "example", "waveform", "spectrum"} else "text"
        metadata = {**(duplicate.metadata_json if duplicate else {}), **payload, "publication_status": "published"}
        if metadata.get("rag_indexed"):
            metadata["index_quality"] = "verified"
        if duplicate:
            unchanged = (
                duplicate.content == text
                and duplicate.course_id == course_id
                and duplicate.material_type == material_type
                and duplicate.topic_id == (topic.id if topic else None)
                and duplicate.metadata_json == metadata
                and duplicate.verified
                and _material_publication_status(duplicate) == "published"
            )
            if unchanged:
                skipped += 1
                continue
            for formula in list(duplicate.formulas):
                db.delete(formula)
            material = duplicate
            updated += 1
        else:
            material = Material(source_document_id=document.id, source_anchor=source_anchor)
            db.add(material)
            published += 1
        material.course_id = course_id
        material.topic_id = topic.id if topic else None
        material.material_type = material_type
        material.title = section or f"导入资料第{item.page_number}页"
        material.content = text
        material.page_number = item.page_number
        material.section = section
        material.tags = [section] if section else []
        material.metadata_json = metadata
        material.verified = True
        db.flush()
        old_chunks = {chunk.chunk_index: chunk for chunk in material.chunks} if duplicate else {}
        for index, text_chunk in enumerate(_chunk_text(text)):
            chunk = old_chunks.pop(index, None)
            if chunk:
                chunk.text = text_chunk
                chunk.keywords = [section] if section else []
            else:
                db.add(MaterialChunk(material_id=material.id, chunk_index=index, text=text_chunk, keywords=[section] if section else []))
        for stale_chunk in old_chunks.values():
            db.execute(update(EvidenceLink).where(EvidenceLink.chunk_id == stale_chunk.id).values(chunk_id=None))
            db.delete(stale_chunk)
        for raw_formula in _formula_sources(payload, verified=True):
            db.add(Formula(material_id=material.id, raw_formula=raw_formula, canonical_expression=raw_formula, variables_json={}))
    job.status = "review" if any(item.review_status == "pending" for item in items) else "published"
    document.status = job.status
    job.summary = {**(job.summary or {}), "published": published, "updated": updated, "skipped": skipped, "revoked": revoked}
    db.commit()
    return {"published": published, "updated": updated, "skipped": skipped, "revoked": revoked}


def _load_ocr_engine() -> Any:
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        try:
            from rapidocr import RapidOCR
        except ImportError as exc:
            raise OCRUnavailable("未安装 rapidocr；请安装 OCR 依赖后再运行显式 OCR") from exc
        _OCR_ENGINE = RapidOCR(params={
            "EngineConfig.onnxruntime.intra_op_num_threads": 2,
            "EngineConfig.onnxruntime.inter_op_num_threads": 1,
            "Global.log_level": "error",
        })
    return _OCR_ENGINE


def _run_ocr(image_bytes: bytes) -> tuple[str, list[dict[str, Any]], float]:
    try:
        import cv2
        import numpy as np
        image = cv2.imdecode(np.frombuffer(image_bytes, dtype=np.uint8), cv2.IMREAD_COLOR)
        result = _load_ocr_engine()(image).to_json() or []
    except OCRUnavailable:
        raise
    except Exception as exc:
        raise OCRUnavailable(f"OCR 执行失败：{exc.__class__.__name__}") from exc
    blocks: list[dict[str, Any]] = []
    for row in result:
        box = row["box"]
        text = row["txt"].strip()
        score = float(row["score"])
        if text:
            blocks.append({"bbox": box, "text": text, "confidence": max(0.0, min(1.0, score))})

    return "\n".join(block["text"] for block in blocks), blocks, (sum(block["confidence"] for block in blocks) / len(blocks) if blocks else 0.0)

def ocr_page(source: SourceDocument | Path | str, page_number: int) -> dict[str, Any]:
    """OCR one page on demand; native ingestion never invokes this implicitly."""
    image = render_page_bytes(source, page_number)
    text, blocks, confidence = _run_ocr(image)
    bounds = _safe_json(_page_bounds_for_source(source, page_number))
    return {"page_number": page_number, "raw_text": text, "normalized_text": text, "blocks": blocks, "boxes": blocks,
            "confidence": confidence, "source_bounds": bounds, "page_bounds": bounds,
            "warning": "OCR 结果必须人工审核" if text else "OCR 未识别到文本"}


def _page_bounds_for_source(source: SourceDocument | Path | str, page_number: int) -> list[float]:
    import fitz
    path = Path(source.path if isinstance(source, SourceDocument) else source)
    with fitz.open(path) as document:
        if page_number < 1 or page_number > len(document):
            raise ValueError("页码超出原始 PDF 范围")
        return _page_bounds(document[page_number - 1])


def ocr_item(db: Session, item_id: int) -> ExtractedItem:
    item = db.get(ExtractedItem, item_id)
    if not item:
        raise LookupError("解析项不存在")
    job = db.get(IngestionJob, item.ingestion_job_id)
    document = db.get(SourceDocument, job.source_document_id) if job else None
    if not document:
        raise LookupError("解析项缺少原始教材")
    result = ocr_page(document, item.page_number)
    payload = dict(item.payload or {})
    payload.update({"ocr_engine": "rapidocr", "ocr_text": result["raw_text"], "ocr_blocks": result["blocks"], "ocr_boxes": result["boxes"], "source_bounds": result["source_bounds"], "page_bounds": result["page_bounds"], "needs_ocr": False, "ocr_warning": result["warning"]})
    payload["quality_review_cleared"] = False
    item.normalized_text = result["normalized_text"]
    item.payload = payload
    item.item_type = _classify(item.normalized_text) if item.normalized_text else "page"
    item.confidence = result["confidence"]
    item.needs_review = True
    item.review_status = "pending"
    _invalidate_item_publication(db, item, "source_item_ocr_pending_review")
    db.commit()
    db.refresh(item)
    return item


def ocr_job_pages(db: Session, job_id: int, pages: list[int] | tuple[int, ...]) -> list[ExtractedItem]:
    """OCR only the requested pages; no background/eager whole-book OCR is performed."""
    requested = sorted({int(page) for page in pages})
    if not requested:
        return []
    items = {item.page_number: item for item in list_job_items(db, job_id)}
    output: list[ExtractedItem] = []
    for page in requested:
        if page not in items:
            raise ValueError(f"任务中不存在第 {page} 页")
        output.append(ocr_item(db, items[page].id))
    return output
